import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).absolute().parents[2]
SCHEMA_PATH = PROJECT_ROOT / "docs" / "silver" / "silver_core_schema.csv"
OUTPUT_PATH = PROJECT_ROOT / "docs" / "data_dictionary.xlsx"

DOCUMENT_STATUS = "IMPLEMENTED_AND_VALIDATED"
SCHEMA_VERSION = "silver_listing_core_v1"
EXPECTED_SCHEMA_FIELDS = 58

MAPPING_COLUMNS = [
    "bronze_column",
    "bronze_type",
    "silver_output_columns",
    "transformation",
    "audit_note",
]

MAPPING_ROWS = [
    ("name", "string", "title", "NULLIF(TRIM(name), '')", "Raw value remains in immutable Bronze and participates in listing_id hash"),
    ("description", "string", "description", "Trim outer whitespace; blank becomes NULL", "Raw value remains in Bronze and participates in hash"),
    ("property_type_name", "string", "property_type_raw; property_type", "Preserve raw; trim/map to five reviewed SRC01 categories", "A new category is schema drift and must stop for review"),
    ("province_name", "string", "province_name_raw; province_name", "Preserve raw; trim and convert blank to NULL", "No administrative-code mapping in Silver Core"),
    ("district_name", "string", "district_name_raw; district_name", "Preserve raw; trim and convert blank to NULL", "No administrative-code mapping in Silver Core"),
    ("ward_name", "string", "ward_name_raw; ward_name", "Preserve raw; trim and convert blank to NULL", "No administrative-code mapping in Silver Core"),
    ("street_name", "string", "street_name", "Trim and convert blank to NULL", "Raw value remains in Bronze and participates in hash"),
    ("project_name", "string", "project_name", "Trim and convert blank to NULL", "Raw value remains in Bronze and participates in hash"),
    ("price", "string", "price_raw; price_vnd; price_per_m2", "Preserve raw; TRY_CAST to DECIMAL(20,0); derive price/m2 only when price > 0 and area > 0", "Listing/asking price in VND, not a completed transaction price"),
    ("area", "double", "area_m2; price_per_m2", "Rename; retain source value; validate positivity; derive price/m2 only when valid", "P99 area is REVIEW rather than deletion"),
    ("floor_count", "double", "floor_count", "Validate nonnegative integer precondition then cast INT", "P99 threshold is property-type-specific"),
    ("frontage_width", "double", "frontage_width_m", "Rename without unit conversion", "Raw value remains available in Bronze"),
    ("house_depth", "double", "house_depth_m", "Rename without unit conversion", "Raw value remains available in Bronze"),
    ("road_width", "double", "road_width_m", "Rename without unit conversion", "Raw value remains available in Bronze"),
    ("bedroom_count", "double", "bedroom_count", "Validate nonnegative integer precondition then cast INT", "P99 threshold is property-type-specific"),
    ("bathroom_count", "double", "bathroom_count", "Validate nonnegative integer precondition then cast INT", "P99 threshold is property-type-specific"),
    ("house_direction", "string", "house_direction", "Trim and convert blank to NULL", "Raw value remains in Bronze and participates in hash"),
    ("balcony_direction", "string", "balcony_direction", "Trim and convert blank to NULL", "Raw value remains in Bronze and participates in hash"),
    ("published_at", "string", "published_at_raw; published_at; published_date; published_year; published_month; year_month; partition_year_month", "Preserve raw; parse TIMESTAMP_NTZ; derive calendar fields without timezone conversion", "Timezone is an open approval decision"),
    ("source_name", "string", "source_name", "Preserve exact technical value vduydong/vietnam-real-estates-2", "Included in listing_id and duplicate_group_id hash"),
    ("source_file", "string", "source_file", "Preserve", "Used as the available duplicate ranking preference"),
    ("ingested_at", "timestamp", "bronze_ingested_at", "Rename; preserve UTC technical timestamp", "Bronze lineage"),
    ("ingestion_date", "date", "bronze_ingestion_date", "Rename", "Bronze lineage"),
    ("batch_id", "string", "bronze_batch_id", "Rename", "Bronze lineage"),
]

DOCUMENT_CONTROL = [
    ("document_status", DOCUMENT_STATUS),
    ("schema_version", SCHEMA_VERSION),
    ("source_scope", "SRC01 only"),
    ("bronze_input_columns", 24),
    ("silver_output_columns", EXPECTED_SCHEMA_FIELDS),
    ("schema_source", "docs/silver/silver_core_schema.csv"),
    ("design_source", "docs/silver/silver_core_design_proposal.md"),
    ("implementation_status", "Implemented by src/silver/01_build_silver_core.py and independently validated"),
]


def read_schema():
    if not SCHEMA_PATH.is_file():
        raise FileNotFoundError(f"Missing Silver schema contract: {SCHEMA_PATH}")
    with SCHEMA_PATH.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    if len(rows) != EXPECTED_SCHEMA_FIELDS:
        raise RuntimeError(
            f"Expected {EXPECTED_SCHEMA_FIELDS} schema fields, found {len(rows)}"
        )
    ordinals = [int(row["ordinal"]) for row in rows]
    column_names = [row["column_name"] for row in rows]
    if ordinals != list(range(1, EXPECTED_SCHEMA_FIELDS + 1)):
        raise RuntimeError("Silver schema ordinals are not contiguous")
    if len(set(column_names)) != EXPECTED_SCHEMA_FIELDS:
        raise RuntimeError("Silver schema contains duplicate column names")
    return rows


def style_sheet(sheet):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="B7B7B7")

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows():
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 30
    for index in range(1, sheet.max_column + 1):
        max_length = max(
            len(str(sheet.cell(row=row, column=index).value or ""))
            for row in range(1, min(sheet.max_row, 60) + 1)
        )
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(max_length + 2, 12), 65
        )


def main():
    schema_rows = read_schema()
    workbook = Workbook()

    schema_sheet = workbook.active
    schema_sheet.title = "silver_core_schema"
    schema_columns = list(schema_rows[0])
    schema_sheet.append(schema_columns)
    for row in schema_rows:
        values = [row[column] for column in schema_columns]
        values[0] = int(values[0])
        schema_sheet.append(values)

    mapping_sheet = workbook.create_sheet("bronze_to_silver")
    mapping_sheet.append(MAPPING_COLUMNS)
    for row in MAPPING_ROWS:
        mapping_sheet.append(row)

    control_sheet = workbook.create_sheet("document_control")
    control_sheet.append(["item", "value"])
    for row in DOCUMENT_CONTROL:
        control_sheet.append(row)

    for sheet in workbook.worksheets:
        style_sheet(sheet)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT_PATH)

    check_workbook = load_workbook(OUTPUT_PATH, read_only=True, data_only=True)
    try:
        check_schema = check_workbook["silver_core_schema"]
        check_mapping = check_workbook["bronze_to_silver"]
        if check_schema.max_row != EXPECTED_SCHEMA_FIELDS + 1:
            raise RuntimeError("Data Dictionary schema read-back validation failed")
        if check_mapping.max_row != len(MAPPING_ROWS) + 1:
            raise RuntimeError("Data Dictionary mapping read-back validation failed")
        if (
            check_workbook["document_control"].cell(row=2, column=2).value
            != DOCUMENT_STATUS
        ):
            raise RuntimeError("Data Dictionary status read-back validation failed")
    finally:
        check_workbook.close()

    print("DATA DICTIONARY: PASS")
    print(f"Output: {OUTPUT_PATH}")
    print(f"Silver fields: {len(schema_rows)}")
    print(f"Bronze mappings: {len(MAPPING_ROWS)}")
    print(f"Status: {DOCUMENT_STATUS}")


if __name__ == "__main__":
    main()
