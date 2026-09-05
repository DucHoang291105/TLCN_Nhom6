import csv
import math
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).absolute().parents[2]
PROFILING_DIR = PROJECT_ROOT / "docs" / "profiling"
THRESHOLDS_PATH = PROFILING_DIR / "silver_dq_thresholds.csv"
COUNT_STATS_PATH = PROFILING_DIR / "count_field_stats.csv"
EXPECTED_HITS_PATH = PROFILING_DIR / "silver_dq_expected_hits.csv"
OUTPUT_PATH = PROJECT_ROOT / "docs" / "data_quality_rules.xlsx"

SOURCE_ID = "SRC01"
RULE_VERSION = "src01_silver_core_dq_v1"
THRESHOLD_VERSION = "src01_raw_20260904_p99_v1"
DOCUMENT_STATUS = "IMPLEMENTED_AND_VALIDATED"

RULE_COLUMNS = [
    "rule_id",
    "field",
    "condition",
    "description",
    "dq_status",
    "field_validity_effect",
    "action",
    "affected_analytics",
    "threshold_scope",
    "rule_version",
    "implementation_notes",
]

RULES = [
    {
        "rule_id": "DQ01",
        "field": "price_raw / price_vnd",
        "condition": "price_raw IS NULL OR TRIM(price_raw) = ''",
        "description": "Thiếu giá đăng",
        "dq_status": "REVIEW",
        "field_validity_effect": "is_price_valid = false",
        "action": "Giữ trong Silver Audit; không dùng cho KPI giá",
        "affected_analytics": "price, price_per_m2",
        "threshold_scope": "NONE",
        "rule_version": RULE_VERSION,
        "implementation_notes": "Không loại khỏi KPI đếm tin đăng nếu KPI đó không cần giá.",
    },
    {
        "rule_id": "DQ02",
        "field": "price_raw / price_vnd",
        "condition": "price_vnd <= 0",
        "description": "Giá đăng không dương",
        "dq_status": "REJECTED",
        "field_validity_effect": "is_price_valid = false",
        "action": "Giữ trong Silver Audit; loại khỏi KPI giá",
        "affected_analytics": "price, price_per_m2",
        "threshold_scope": "NON_POSITIVE",
        "rule_version": RULE_VERSION,
        "implementation_notes": "REJECTED chỉ áp dụng cho mục đích dùng giá, không phải xóa toàn bộ record.",
    },
    {
        "rule_id": "DQ03",
        "field": "price_raw / price_vnd",
        "condition": "price_raw IS NOT NULL AND TRIM(price_raw) <> '' AND TRY_CAST(price_raw AS DECIMAL(20,0)) IS NULL",
        "description": "Giá có dữ liệu nhưng không chuyển được sang số",
        "dq_status": "REJECTED",
        "field_validity_effect": "is_price_valid = false",
        "action": "Giữ price_raw, đặt price_vnd = NULL; loại khỏi KPI giá",
        "affected_analytics": "price, price_per_m2",
        "threshold_scope": "CAST",
        "rule_version": RULE_VERSION,
        "implementation_notes": "Điều kiện loại trừ NULL/blank nên không còn bắt trùng DQ01.",
    },
    {
        "rule_id": "DQ04",
        "field": "area_m2",
        "condition": "area_m2 IS NULL OR area_m2 <= 0",
        "description": "Diện tích thiếu hoặc không dương",
        "dq_status": "REJECTED",
        "field_validity_effect": "is_area_valid = false",
        "action": "Giữ trong Silver Audit; không tính price_per_m2",
        "affected_analytics": "area, price_per_m2",
        "threshold_scope": "POSITIVE",
        "rule_version": RULE_VERSION,
        "implementation_notes": "Record vẫn có thể dùng cho KPI không phụ thuộc diện tích.",
    },
    {
        "rule_id": "DQ05",
        "field": "area_m2",
        "condition": "area_m2 > COALESCE(property_type_threshold.p99_area_m2, global_fallback.p99_area_m2)",
        "description": "Diện tích lớn hơn P99 của cùng loại bất động sản",
        "dq_status": "REVIEW",
        "field_validity_effect": "Không tự động đổi is_area_valid",
        "action": "Giữ giá trị và gắn cờ review; Gold sạch tự chọn có loại outlier hay không",
        "affected_analytics": "area distribution, price_per_m2",
        "threshold_scope": "P99_BY_PROPERTY_TYPE_WITH_GLOBAL_FALLBACK",
        "rule_version": RULE_VERSION,
        "implementation_notes": "P99 area được tính trên mọi area > 0, không phụ thuộc price.",
    },
    {
        "rule_id": "DQ06",
        "field": "price_per_m2",
        "condition": "unrounded(price_vnd / area_m2) > COALESCE(property_type_threshold.p99_price_per_m2, global_fallback.p99_price_per_m2)",
        "description": "Giá đăng trên m² lớn hơn P99 của cùng loại bất động sản",
        "dq_status": "REVIEW",
        "field_validity_effect": "Không tự động đổi is_price_valid/is_area_valid",
        "action": "Giữ giá trị và gắn cờ review; mặc định loại khỏi KPI giá sạch",
        "affected_analytics": "price_per_m2",
        "threshold_scope": "P99_BY_PROPERTY_TYPE_WITH_GLOBAL_FALLBACK",
        "rule_version": RULE_VERSION,
        "implementation_notes": "So P99 bằng tỷ lệ chưa làm tròn như profiling; cột output price_per_m2 được lưu DECIMAL(24,2). Chỉ tính khi price > 0 và area > 0.",
    },
    {
        "rule_id": "DQ07",
        "field": "province_name",
        "condition": "province_name IS NULL OR TRIM(province_name) = ''",
        "description": "Thiếu tỉnh/thành phố nguồn",
        "dq_status": "REJECTED",
        "field_validity_effect": "is_province_valid = false",
        "action": "Giữ trong Silver Audit; không dùng cho KPI địa lý",
        "affected_analytics": "province, district, ward geography",
        "threshold_scope": "NONE",
        "rule_version": RULE_VERSION,
        "implementation_notes": "Mapping mã hành chính thuộc Silver Enriched, không thuộc Core.",
    },
    {
        "rule_id": "DQ08",
        "field": "district_name",
        "condition": "district_name IS NULL OR TRIM(district_name) = ''",
        "description": "Thiếu quận/huyện nguồn",
        "dq_status": "REVIEW",
        "field_validity_effect": "is_district_valid = false",
        "action": "Giữ record; vẫn có thể dùng cho KPI cấp tỉnh",
        "affected_analytics": "district, ward geography",
        "threshold_scope": "NONE",
        "rule_version": RULE_VERSION,
        "implementation_notes": "Không dùng row status toàn cục để loại khỏi KPI cấp tỉnh.",
    },
    {
        "rule_id": "DQ09",
        "field": "ward_name",
        "condition": "ward_name IS NULL OR TRIM(ward_name) = ''",
        "description": "Thiếu phường/xã nguồn",
        "dq_status": "REVIEW",
        "field_validity_effect": "is_ward_valid = false",
        "action": "Giữ record; vẫn có thể dùng cho KPI cấp tỉnh/quận",
        "affected_analytics": "ward geography",
        "threshold_scope": "NONE",
        "rule_version": RULE_VERSION,
        "implementation_notes": "Không dùng row status toàn cục để loại khỏi KPI cấp trên.",
    },
    {
        "rule_id": "DQ10",
        "field": "published_at_raw / published_at",
        "condition": "published_at_raw IS NULL OR TRIM(published_at_raw) = '' OR TRY_CAST(published_at_raw AS TIMESTAMP_NTZ) IS NULL",
        "description": "Ngày đăng thiếu hoặc không parse được",
        "dq_status": "REJECTED",
        "field_validity_effect": "is_published_at_valid = false",
        "action": "Giữ published_at_raw, đặt published_at = NULL; không dùng cho KPI thời gian",
        "affected_analytics": "date, month, year trends",
        "threshold_scope": "CAST",
        "rule_version": RULE_VERSION,
        "implementation_notes": "Raw không có UTC offset; Silver v1 lưu TIMESTAMP_NTZ/source-local-naive và không tự gán UTC.",
    },
    {
        "rule_id": "DQ11",
        "field": "record",
        "condition": "duplicate_rank > 1 within duplicate_group_id built from source_id + source_name + all 19 raw business fields",
        "description": "Bản sao nội nguồn của cùng nhóm exact duplicate",
        "dq_status": "REJECTED",
        "field_validity_effect": "is_canonical = false",
        "action": "Giữ mọi row trong Silver Audit; chỉ duplicate_rank = 1 vào analytical/canonical view",
        "affected_analytics": "all aggregate counts and measures when canonical rows are required",
        "threshold_scope": "EXACT_INTRA_SOURCE_DUPLICATE",
        "rule_version": RULE_VERSION,
        "implementation_notes": "Dùng SHA-256 trên JSON null-safe có version của source_id, source_name và đủ 19 raw fields; không trim/cast trước hash. Baseline full-row: 36 nhóm, 50 dòng dư. Các bản sao exact trong cùng file không có physical row ID ổn định; canonical business row vẫn tương đương.",
    },
    {
        "rule_id": "DQ12",
        "field": "floor_count",
        "condition": "floor_count IS NOT NULL AND floor_count > COALESCE(property_type_threshold.p99_floor_count, 8)",
        "description": "Số tầng lớn hơn P99 theo loại BĐS",
        "dq_status": "REVIEW",
        "field_validity_effect": "Không tự động đổi tính hợp lệ của record",
        "action": "Giữ record và giá trị; gắn cờ để review",
        "affected_analytics": "floor_count",
        "threshold_scope": "P99_BY_PROPERTY_TYPE; GLOBAL_FALLBACK=8",
        "rule_version": RULE_VERSION,
        "implementation_notes": "Dùng quantile_disc và toán tử >. SRC01 hiện có 0 giá trị âm/không nguyên; validator phải dừng nếu precondition này thay đổi trước khi cast INT.",
    },
    {
        "rule_id": "DQ13",
        "field": "bedroom_count",
        "condition": "bedroom_count IS NOT NULL AND bedroom_count > COALESCE(property_type_threshold.p99_bedroom_count, 14)",
        "description": "Số phòng ngủ lớn hơn P99 theo loại BĐS",
        "dq_status": "REVIEW",
        "field_validity_effect": "Không tự động đổi tính hợp lệ của record",
        "action": "Giữ record và giá trị; gắn cờ để review",
        "affected_analytics": "bedroom_count",
        "threshold_scope": "P99_BY_PROPERTY_TYPE; GLOBAL_FALLBACK=14",
        "rule_version": RULE_VERSION,
        "implementation_notes": "Dùng quantile_disc và toán tử >. SRC01 hiện có 0 giá trị âm/không nguyên; validator phải dừng nếu precondition này thay đổi trước khi cast INT.",
    },
    {
        "rule_id": "DQ14",
        "field": "bathroom_count",
        "condition": "bathroom_count IS NOT NULL AND bathroom_count > COALESCE(property_type_threshold.p99_bathroom_count, 13)",
        "description": "Số phòng tắm lớn hơn P99 theo loại BĐS",
        "dq_status": "REVIEW",
        "field_validity_effect": "Không tự động đổi tính hợp lệ của record",
        "action": "Giữ record và giá trị; gắn cờ để review",
        "affected_analytics": "bathroom_count",
        "threshold_scope": "P99_BY_PROPERTY_TYPE; GLOBAL_FALLBACK=13",
        "rule_version": RULE_VERSION,
        "implementation_notes": "Dùng quantile_disc và toán tử >. SRC01 hiện có 0 giá trị âm/không nguyên; validator phải dừng nếu precondition này thay đổi trước khi cast INT.",
    },
]

PROFILING_BASELINE = [
    ("total_rows", 3_500_744, "Tổng record SRC01"),
    ("null_price", 218_494, "Thiếu giá"),
    ("zero_price", 57_320, "Giá bằng 0"),
    ("negative_price", 0, "Giá nhỏ hơn 0"),
    ("invalid_non_null_price", 0, "Giá có dữ liệu nhưng không cast được"),
    ("null_area", 4, "Thiếu diện tích"),
    ("blank_name", 1_868, "Tên rỗng sau TRIM; đang là quyết định DQ mở"),
    ("blank_description", 1_868, "Mô tả rỗng sau TRIM; đang là quyết định DQ mở"),
    ("exact_duplicate_groups", 36, "Nhóm exact duplicate theo DQ11"),
    ("exact_duplicate_excess_rows", 50, "Số bản sao dư theo DQ11"),
    ("published_at_with_utc_offset", 0, "Không quan sát thấy UTC offset trong raw timestamp"),
]

DESIGN_NOTES = [
    ("document_status", DOCUMENT_STATUS),
    ("source_scope", "Chỉ SRC01 / Silver Listing Core; chưa bao gồm mapping location"),
    ("condition_dialect", "Điều kiện trong workbook là Spark SQL-style pseudocode; code thực tế phải dùng PySpark expressions/type contract tương đương"),
    ("retention", "Mọi 3,500,744 row được giữ trong Silver Audit; không drop âm thầm"),
    ("status_precedence", "Nếu có REJECTED => dq_status=REJECTED; nếu không nhưng có REVIEW => REVIEW; còn lại VALID"),
    ("gold_filtering", "Gold phải lọc theo is_canonical và field-validity/error-code liên quan KPI; không dùng một filter dq_status toàn cục"),
    ("threshold_policy", "DQ05/DQ06/DQ12-DQ14 dùng P99 theo property_type; chỉ dùng global fallback khi loại BĐS chưa có threshold"),
    ("threshold_execution", "Pipeline phải đọc đúng threshold version từ CSV; không tự tính lại bằng Spark vì thuật toán percentile có thể khác"),
    ("outlier_meaning", "Vượt P99 là REVIEW, không chứng minh dữ liệu sai và không tự động null/drop"),
    ("duplicate_policy", "DQ11 chỉ dedup exact duplicate nội SRC01. Cross-source entity matching là bài toán riêng trong tương lai"),
    ("timezone_contract", "published_at không có offset; Silver v1 dùng TIMESTAMP_NTZ/source-local-naive và không tự gán UTC"),
    ("candidate_dq15", "DEFERRED_NOT_IMPLEMENTED: 1,868 row có name và description rỗng sau TRIM; DQ v1 chỉ gồm DQ01-DQ14"),
    ("implementation_status", "IMPLEMENTED_AND_VALIDATED: src/silver/01_build_silver_core.py + 02_validate_silver_core.py"),
]


def read_csv_rows(path):
    if not path.is_file():
        raise FileNotFoundError(
            f"Missing {path}. Run src/quality/05_profile_count_fields.py first."
        )
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def parse_cell(value):
    if value is None or value == "":
        return None
    try:
        number = float(value)
    except ValueError:
        return value
    return int(number) if number.is_integer() else number


def validate_inputs(threshold_rows, count_stat_rows, expected_hit_rows):
    for label, rows in (
        ("thresholds", threshold_rows),
        ("count stats", count_stat_rows),
        ("expected hits", expected_hit_rows),
    ):
        source_ids = {row["source_id"] for row in rows}
        versions = {row["threshold_version"] for row in rows}
        if source_ids != {SOURCE_ID}:
            raise RuntimeError(f"Unexpected source IDs in {label}: {source_ids}")
        if versions != {THRESHOLD_VERSION}:
            raise RuntimeError(f"Unexpected threshold versions in {label}: {versions}")

    expected_scopes = {
        "__GLOBAL_FALLBACK__": "GLOBAL_FALLBACK",
        "Biệt thự/Nhà liền kề": "PROPERTY_TYPE",
        "Căn hộ chung cư": "PROPERTY_TYPE",
        "Nhà": "PROPERTY_TYPE",
        "Shophouse": "PROPERTY_TYPE",
        "Đất": "PROPERTY_TYPE",
    }
    actual_scopes = {
        row["property_type"]: row["threshold_scope"] for row in threshold_rows
    }
    if actual_scopes != expected_scopes:
        raise RuntimeError(f"Unexpected threshold scopes: {actual_scopes}")

    expected_count_p99 = {
        "__GLOBAL_FALLBACK__": (8, 14, 13),
        "Biệt thự/Nhà liền kề": (7, 10, 10),
        "Căn hộ chung cư": (40, 4, 4),
        "Nhà": (8, 19, 18),
        "Shophouse": (7, 10, 10),
        "Đất": (9, 20, 15),
    }
    actual_count_p99 = {
        row["property_type"]: (
            int(float(row["p99_floor_count"])),
            int(float(row["p99_bedroom_count"])),
            int(float(row["p99_bathroom_count"])),
        )
        for row in threshold_rows
    }
    if actual_count_p99 != expected_count_p99:
        raise RuntimeError(f"Unexpected count thresholds: {actual_count_p99}")

    expected_continuous_p99 = {
        "__GLOBAL_FALLBACK__": (2_222.0, 607_692_307.6923077),
        "Biệt thự/Nhà liền kề": (894.2520000000252, 625_000_000.0),
        "Căn hộ chung cư": (240.0, 261_864_406.77966103),
        "Nhà": (544.0, 735_000_000.0),
        "Shophouse": (535.0, 552_310_374.89102),
        "Đất": (10_716.349999999977, 347_947_879.1170806),
    }
    for row in threshold_rows:
        expected_area, expected_price_m2 = expected_continuous_p99[
            row["property_type"]
        ]
        actual_area = float(row["p99_area_m2"])
        actual_price_m2 = float(row["p99_price_per_m2"])
        if not math.isclose(actual_area, expected_area, rel_tol=1e-12, abs_tol=1e-6):
            raise RuntimeError(
                f"Unexpected P99 area for {row['property_type']}: {actual_area}"
            )
        if not math.isclose(
            actual_price_m2,
            expected_price_m2,
            rel_tol=1e-12,
            abs_tol=1e-6,
        ):
            raise RuntimeError(
                f"Unexpected P99 price/m2 for {row['property_type']}: "
                f"{actual_price_m2}"
            )

    global_count_rows = [row for row in count_stat_rows if row["scope"] == "ALL"]
    if len(global_count_rows) != 3:
        raise RuntimeError("Count-field global profiling is incomplete")
    if any(int(float(row["negative_rows"])) for row in global_count_rows):
        raise RuntimeError("Negative count values found")
    if any(int(float(row["non_integer_rows"])) for row in global_count_rows):
        raise RuntimeError("Non-integer count values found")

    expected_hits = {
        "DQ05": 34_807,
        "DQ06": 32_219,
        "DQ12": 6_511,
        "DQ13": 14_530,
        "DQ14": 11_636,
    }
    actual_hits = {
        row["rule_id"]: int(float(row["expected_hit_rows"]))
        for row in expected_hit_rows
    }
    if actual_hits != expected_hits:
        raise RuntimeError(f"Unexpected DQ hit baseline: {actual_hits}")
    if {
        (row["threshold_scope"], row["comparison"])
        for row in expected_hit_rows
    } != {("P99_BY_PROPERTY_TYPE", "STRICT_GREATER_THAN")}:
        raise RuntimeError("Unexpected expected-hit threshold contract")


def append_dict_sheet(workbook, title, columns, rows):
    sheet = workbook.create_sheet(title)
    sheet.append(columns)
    for row in rows:
        sheet.append([parse_cell(row.get(column)) for column in columns])
    return sheet


def style_sheet(sheet):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="B7B7B7")

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows():
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 30
    for index in range(1, sheet.max_column + 1):
        max_length = max(
            len(str(sheet.cell(row=row, column=index).value or ""))
            for row in range(1, min(sheet.max_row, 50) + 1)
        )
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(max_length + 2, 12), 65
        )


def main():
    threshold_rows = read_csv_rows(THRESHOLDS_PATH)
    count_stat_rows = read_csv_rows(COUNT_STATS_PATH)
    expected_hit_rows = read_csv_rows(EXPECTED_HITS_PATH)
    validate_inputs(threshold_rows, count_stat_rows, expected_hit_rows)

    workbook = Workbook()
    rules_sheet = workbook.active
    rules_sheet.title = "dq_rules"
    rules_sheet.append(RULE_COLUMNS)
    for rule in RULES:
        rules_sheet.append([rule[column] for column in RULE_COLUMNS])

    threshold_columns = list(threshold_rows[0])
    append_dict_sheet(
        workbook,
        "property_type_thresholds",
        threshold_columns,
        threshold_rows,
    )

    count_columns = list(count_stat_rows[0])
    append_dict_sheet(
        workbook,
        "count_field_profiling",
        count_columns,
        count_stat_rows,
    )

    hit_columns = list(expected_hit_rows[0])
    append_dict_sheet(
        workbook,
        "expected_rule_hits",
        hit_columns,
        expected_hit_rows,
    )

    baseline_sheet = workbook.create_sheet("profiling_baseline")
    baseline_sheet.append(["metric", "value", "description"])
    for row in PROFILING_BASELINE:
        baseline_sheet.append(row)

    notes_sheet = workbook.create_sheet("design_notes")
    notes_sheet.append(["item", "decision_or_note"])
    for row in DESIGN_NOTES:
        notes_sheet.append(row)

    for sheet in workbook.worksheets:
        style_sheet(sheet)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT_PATH)

    check_workbook = load_workbook(OUTPUT_PATH, read_only=True, data_only=True)
    try:
        check_rules = check_workbook["dq_rules"]
        ids = [
            check_rules.cell(row=row, column=1).value
            for row in range(2, check_rules.max_row + 1)
        ]
        if ids != [f"DQ{number:02d}" for number in range(1, 15)]:
            raise RuntimeError("DQ workbook rule read-back validation failed")
        if check_workbook["property_type_thresholds"].max_row != 7:
            raise RuntimeError("DQ workbook threshold read-back validation failed")
    finally:
        check_workbook.close()

    print("DATA QUALITY RULES: PASS")
    print(f"Output: {OUTPUT_PATH}")
    print(f"Rules: {len(RULES)}")
    print(f"Rule version: {RULE_VERSION}")
    print(f"Threshold version: {THRESHOLD_VERSION}")
    print(f"Status: {DOCUMENT_STATUS}")


if __name__ == "__main__":
    main()
