from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).absolute().parents[2]
OUTPUT_PATH = PROJECT_ROOT / "docs" / "data_sources.xlsx"
CATALOG_VERSION = "1.0"
LAST_REVIEWED = "2026-09-05"

COLUMNS = [
    "source_id",
    "source_name",
    "source_display_name",
    "url",
    "canonical_url",
    "source_type",
    "grain",
    "role",
    "bronze_output",
    "silver_output",
    "owner",
    "current_status",
    "scope",
    "repository_owner",
    "publisher",
    "license",
    "format",
    "update_mode",
    "natural_key",
    "timezone",
    "currency",
    "administrative_vintage",
    "snapshot_revision",
    "snapshot_date",
    "download_status",
    "checksum_sha256",
    "schema_version",
    "notes",
]

SOURCES = [
    {
        "source_id": "SRC01",
        "source_name": "vduydong/vietnam-real-estates-2",
        "source_display_name": "Vietnam Real Estate Listings 2025-2026",
        "url": "https://huggingface.co/datasets/vduydong/vietnam-real-estates-2",
        "canonical_url": "https://huggingface.co/datasets/vduydong/vietnam-real-estates-2",
        "source_type": "listing",
        "grain": "1 record = 1 real-estate listing",
        "role": "Main fact source for the TLCN real-estate pipeline",
        "bronze_output": "data/bronze/real_estate/",
        "silver_output": "data/silver/real_estate_core/",
        "owner": "Listing Pipeline",
        "current_status": "ACTIVE - BRONZE PASS; SILVER CORE PASS",
        "scope": "TLCN_CONFIRMED",
        "repository_owner": "vduydong",
        "publisher": "vduydong repository; dataset card cites TiniX AI",
        "license": "CC BY-NC 4.0",
        "format": "Parquet",
        "update_mode": "Local snapshot / batch",
        "natural_key": "No native listing ID in local schema; versioned source-scoped hash required",
        "timezone": "published_at has no UTC offset; Silver v1 preserves source-local TIMESTAMP_NTZ",
        "currency": "VND",
        "administrative_vintage": "63-province listing labels; preserve source names",
        "snapshot_revision": "HF 0f1e5586850507e2a1d0ce6d069424d8e6ffd263; local: 10 shards, 3,500,744 rows, 19 raw columns",
        "snapshot_date": "Local snapshot verified 2026-09-04",
        "download_status": "DOWNLOADED_LOCAL",
        "checksum_sha256": "Not recorded for initial snapshot",
        "schema_version": "src01_raw_v1",
        "notes": "Silver Core v1 implemented and independently validated. Location enrichment remains a separate later step.",
    },
    {
        "source_id": "SRC02",
        "source_name": "adminvsrm/GISData",
        "source_display_name": "Vietnam GIS Data",
        "url": "https://github.com/adminvsrm/GISData",
        "canonical_url": "https://github.com/nguyenduy1133/Free-GIS-Data",
        "source_type": "GIS / geographic reference",
        "grain": "1 geographic feature (province polygon, headquarters point, or file-specific boundary)",
        "role": "Geometry for maps and province-level checks; supporting input to Location Master",
        "bronze_output": "data/bronze/geography/gisdata/ (planned)",
        "silver_output": "silver_geographic_boundary + silver_geographic_point + bridge_admin_geometry (planned)",
        "owner": "Location Mapping",
        "current_status": "CONFIRMED - SOURCE PROFILING ASSIGNED",
        "scope": "TLCN_CONFIRMED",
        "repository_owner": "adminvsrm (fork)",
        "publisher": "Nguyễn Duy Liêm (nguyenduy1133 upstream author)",
        "license": "FORMAL LICENSE NOT SPECIFIED; README states public use with attribution",
        "format": "GeoJSON and Shapefile (SHP)",
        "update_mode": "Versioned repository snapshot",
        "natural_key": "Must be established per selected file/schema",
        "timezone": "N/A",
        "currency": "N/A",
        "administrative_vintage": "Contains pre-2025 63-province and post-2025 34-province assets",
        "snapshot_revision": "Fork 7645534d0a482ee867f26f137d3dd3fc54d9446f; upstream ccb9f4ae992418bfeefd06da1eb42d0249632176; pin one before ingest",
        "snapshot_date": "Fork 2025-07-10; upstream 2025-07-23",
        "download_status": "NOT_DOWNLOADED_BY_PERSON_1",
        "checksum_sha256": "Pending source selection",
        "schema_version": "Pending source profiling",
        "notes": "Do not union with listings. Coverage is not nationwide district/ward geometry: headquarters files may be southwest-only and ward boundaries include a Thu Duc subset. Select exact files, commit and vintage before use.",
    },
    {
        "source_id": "SRC03",
        "source_name": "thanglequoc/vietnamese-provinces-database",
        "source_display_name": "Vietnamese Provinces Database",
        "url": "https://github.com/thanglequoc/vietnamese-provinces-database",
        "canonical_url": "https://github.com/thanglequoc/vietnamese-provinces-database",
        "source_type": "administrative reference / master data",
        "grain": "1 row = 1 region, unit type, province, district (historical versions), or ward/commune, depending on table",
        "role": "Community-maintained administrative names/codes/hierarchy for versioned Location Master / Dim_Location",
        "bronze_output": "data/bronze/administrative/vietnamese_provinces/ (planned)",
        "silver_output": "silver_admin_unit_versioned + silver_admin_alias + optional silver_admin_geometry (planned)",
        "owner": "Location Mapping",
        "current_status": "CONFIRMED - SOURCE PROFILING ASSIGNED",
        "scope": "TLCN_CONFIRMED",
        "repository_owner": "thanglequoc",
        "publisher": "Thang Le Quoc; independent community project derived from statistical-office API data",
        "license": "MIT",
        "format": "SQL, non-SQL datasets and GIS add-ons",
        "update_mode": "Versioned releases; choose an administrative vintage",
        "natural_key": "Administrative unit code (verify selected release schema)",
        "timezone": "N/A",
        "currency": "N/A",
        "administrative_vintage": "Versioned: v2.4.1 has 63 provinces/705 districts/10,599 wards; v3.0.2 covers the 34-province two-tier transition; v5.0.0 is a later 34-province/3,321-ward current rollup without districts",
        "snapshot_revision": "Plan: v2.4.1@fc33b7411ec4e3697817fab8118718a8d39ef090 + v3.0.2@78594ea643d5fdebd640ea3889ea91f84b9463ae; evaluate v5.0.0@b092d6b45ea76c39990afd34375eabe1f6c3a492 separately",
        "snapshot_date": "Releases not downloaded; selection pending Location Master work",
        "download_status": "NOT_DOWNLOADED_BY_PERSON_1",
        "checksum_sha256": "Pending source selection",
        "schema_version": "Pending source profiling",
        "notes": "Not a government/NSO official repository. Do not clone blindly or replace historical districts with the latest two-tier schema. Select files/releases, verify them, and combine with SRC02 through a bridge rather than unioning with listings.",
    },
    {
        "source_id": "SRC04",
        "source_name": "tmquan/nso-gov-vn",
        "source_display_name": "Vietnam NSO PX-Web Dataset",
        "url": "https://huggingface.co/datasets/tmquan/nso-gov-vn",
        "canonical_url": "https://huggingface.co/datasets/tmquan/nso-gov-vn",
        "source_type": "socioeconomic / statistical",
        "grain": "1 PX-Web statistical cell in long format, or 1 matrix-specific record",
        "role": "Optional socioeconomic context for Gold/DWH",
        "bronze_output": "data/bronze/socioeconomic/nso/ (future only)",
        "silver_output": "data/silver/socioeconomic/ (future only)",
        "owner": "Chưa phân công",
        "current_status": "PLANNED / OPTIONAL - DO NOT IMPLEMENT",
        "scope": "OPTIONAL_FUTURE",
        "repository_owner": "tmquan",
        "publisher": "TMQuan mirror; underlying data from Vietnam NSO PX-Web",
        "license": "CC BY-NC 4.0 redistribution; verify underlying NSO terms",
        "format": "Parquet; native matrices and unified long format",
        "update_mode": "Versioned mirror snapshot",
        "natural_key": "table_id + dimensions + time (to be validated if activated)",
        "timezone": "Varies by indicator / generally N/A",
        "currency": "Varies by indicator",
        "administrative_vintage": "Varies by matrix and reference year",
        "snapshot_revision": "Available HF revision 0fa139f597ac6c0d78419cb30ca70ee18926dd97; 502 matrices / 316,108 long-format cells / 12 databases reported by dataset card",
        "snapshot_date": "Not downloaded; revision observed 2026-09-04",
        "download_status": "DO_NOT_DOWNLOAD_YET",
        "checksum_sha256": "N/A until activated",
        "schema_version": "Not defined",
        "notes": "Separate Bronze/Silver domain. Never union with listing rows.",
    },
    {
        "source_id": "SRC05",
        "source_name": "google/maps-places-api",
        "source_display_name": "Google Maps / Places API",
        "url": "https://developers.google.com/maps",
        "canonical_url": "https://developers.google.com/maps",
        "source_type": "external geocoding / POI API",
        "grain": "1 API response per geocoding/place request",
        "role": "Possible future geocoding, POI and distance enrichment",
        "bronze_output": "N/A - not approved",
        "silver_output": "N/A - not approved",
        "owner": "Chưa phân công",
        "current_status": "NOT CONFIRMED - DO NOT IMPLEMENT",
        "scope": "DO_NOT_IMPLEMENT",
        "repository_owner": "N/A",
        "publisher": "Google",
        "license": "Google Maps Platform Terms if activated",
        "format": "REST API / JSON",
        "update_mode": "On-demand API",
        "natural_key": "place_id/request key (only if activated)",
        "timezone": "N/A",
        "currency": "N/A",
        "administrative_vintage": "Current service response; not an audit master",
        "snapshot_revision": "N/A",
        "snapshot_date": "N/A",
        "download_status": "DO_NOT_CALL_OR_DOWNLOAD",
        "checksum_sha256": "N/A",
        "schema_version": "Not defined",
        "notes": "Do not make bulk calls for 3.5M listings. Requires explicit approval, key and cost review.",
    },
]


def style_sheet(sheet):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="B7B7B7")

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in sheet.iter_rows():
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 30


def main():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "data_sources"
    sheet.append(COLUMNS)

    for source in SOURCES:
        sheet.append([source[column] for column in COLUMNS])

    for row_number in range(2, len(SOURCES) + 2):
        for url_column in ("url", "canonical_url"):
            url_cell = sheet.cell(
                row=row_number,
                column=COLUMNS.index(url_column) + 1,
            )
            url_cell.hyperlink = url_cell.value
            url_cell.style = "Hyperlink"
        sheet.row_dimensions[row_number].height = 90

    widths = {
        "source_id": 12,
        "source_name": 44,
        "source_display_name": 38,
        "url": 58,
        "canonical_url": 58,
        "source_type": 30,
        "grain": 52,
        "role": 52,
        "bronze_output": 46,
        "silver_output": 58,
        "owner": 18,
        "current_status": 38,
        "scope": 22,
        "repository_owner": 24,
        "publisher": 52,
        "license": 48,
        "format": 34,
        "update_mode": 34,
        "natural_key": 52,
        "timezone": 40,
        "currency": 18,
        "administrative_vintage": 64,
        "snapshot_revision": 80,
        "snapshot_date": 38,
        "download_status": 32,
        "checksum_sha256": 28,
        "schema_version": 28,
        "notes": 80,
    }
    for index, column in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = widths[column]
    style_sheet(sheet)

    scope_sheet = workbook.create_sheet("scope_and_rules")
    scope_sheet.append(["item", "value"])
    scope_rows = [
        ("catalog_version", CATALOG_VERSION),
        ("last_reviewed", LAST_REVIEWED),
        ("confirmed_tlcn_sources", "SRC01, SRC02, SRC03"),
        ("optional_future_source", "SRC04"),
        ("do_not_implement", "SRC05"),
        ("rule_1", "Each source domain has its own Bronze ingestion."),
        ("rule_2", "GIS/Admin/NSO must not be unioned with listing rows."),
        ("rule_3", "Silver Core currently applies only to SRC01."),
        ("rule_4", "SRC02 + SRC03 build a versioned Location Master."),
        ("rule_5", "Only future listing sources map to a canonical listing schema."),
        ("rule_6", "SRC04/SRC05 require explicit approval before implementation."),
    ]
    for row in scope_rows:
        scope_sheet.append(row)
    scope_sheet.column_dimensions["A"].width = 30
    scope_sheet.column_dimensions["B"].width = 100
    style_sheet(scope_sheet)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT_PATH)

    check_workbook = load_workbook(OUTPUT_PATH, read_only=True, data_only=True)
    try:
        check_sheet = check_workbook["data_sources"]
        actual_headers = [cell.value for cell in next(check_sheet.iter_rows())]
        if actual_headers != COLUMNS or check_sheet.max_row != len(SOURCES) + 1:
            raise RuntimeError("data_sources.xlsx read-back validation failed")
    finally:
        check_workbook.close()

    print("DATA SOURCE CATALOG: PASS")
    print(f"Output: {OUTPUT_PATH}")
    print(f"Sources: {len(SOURCES)}")
    print("TLCN confirmed: SRC01, SRC02, SRC03")
    print("Optional/do not implement: SRC04, SRC05")


if __name__ == "__main__":
    main()
